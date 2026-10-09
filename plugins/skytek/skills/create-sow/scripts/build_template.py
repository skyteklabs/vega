# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "resvg-py>=0.2", "Pillow>=10"]
# ///
"""Build the SkyTek SOW template (assets/sow-template.docx).

Section layout follows the Google DAF/PSF partner SOW template, without its
Google funding text. Every docxtpl tag is written as a single run, so Word
never splits it. The logo is rasterised from assets/skytek-logo.svg, since
Word templates built with python-docx cannot embed SVG. The same artwork,
turned 45° at 20% opacity, is anchored behind the text of every page as a watermark.

    uv run build_template.py [--out PATH]
"""

import argparse
import io
import re
from pathlib import Path

import resvg_py
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

WATERMARK_OPACITY = 0.2
WATERMARK_ANGLE = 45  # degrees, counter-clockwise

ASSETS = Path(__file__).resolve().parent.parent / "assets"
DEFAULT_OUT = ASSETS / "sow-template.docx"
LOGO_SVG = ASSETS / "skytek-logo.svg"

# Logo colours: the mark's blue, and the wordmark's ink.
BLUE = RGBColor(0x00, 0x64, 0xE0)
INK = RGBColor(0x1C, 0x2B, 0x33)
MUTED = RGBColor(0x5F, 0x6B, 0x73)
HEADER_FILL = "0064E0"
FONT = "Arial"

SECTIONS = [
    "Executive Summary",
    "Requirements and Solution Overview",
    "Activities",
    "Deliverables",
    "Out of Scope, Assumptions and Risks",
    "Success Criteria",
    "Estimated Timeline",
    "Project Roles",
    "Costs",
    "Acceptance",
]


def logo_png():
    """The SVG minus its white backdrop, cropped to the artwork, as PNG bytes."""
    svg = LOGO_SVG.read_text()
    svg = re.sub(r'<rect [^>]*style="fill:white;"/>', "", svg)
    svg = re.sub(
        r'width="100%" height="100%" viewBox="[^"]+"',
        'width="3633" height="815" viewBox="240 240 3633 815"',
        svg,
    )
    return bytes(resvg_py.svg_to_bytes(svg_string=svg, width=1600))


def watermark_png():
    """The logo at WATERMARK_OPACITY, turned WATERMARK_ANGLE, baked into the pixels.

    Baking alpha and rotation in (rather than OOXML transparency and rotation
    attributes) keeps LibreOffice's headless PDF conversion faithful.
    """
    img = Image.open(io.BytesIO(logo_png())).convert("RGBA")
    img = img.rotate(WATERMARK_ANGLE, expand=True, resample=Image.BICUBIC)
    r, g, b, a = img.split()
    img.putalpha(a.point(lambda v: int(v * WATERMARK_OPACITY)))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def floating(paragraph, image_bytes, width, height, behind=True):
    """Anchor a picture to the page, centered, optionally behind the text."""
    run = paragraph.add_run()
    run.add_picture(io.BytesIO(image_bytes), width=width, height=height)
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))

    anchor = OxmlElement("wp:anchor")
    for attr, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "1", "locked": "0",
        "layoutInCell": "1", "allowOverlap": "1", "behindDoc": "1" if behind else "0",
    }.items():
        anchor.set(attr, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)
    for tag_name in ("positionH", "positionV"):
        pos = OxmlElement(f"wp:{tag_name}")
        pos.set("relativeFrom", "page")
        align = OxmlElement("wp:align")
        align.text = "center"
        pos.append(align)
        anchor.append(pos)
    for child in ("wp:extent", "wp:effectExtent"):
        el = inline.find(qn(child))
        if el is not None:
            anchor.append(el)
    anchor.append(OxmlElement("wp:wrapNone"))
    for child in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        el = inline.find(qn(child))
        if el is not None:
            anchor.append(el)

    drawing.remove(inline)
    drawing.append(anchor)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def field(paragraph, instr, size, color):
    """A Word field such as PAGE at the end of `paragraph`; Word and LibreOffice fill it in.

    Each part of the field is its own sized run, so the page number matches the text.
    """
    for part in ("begin", instr, "separate", "1", "end"):
        run = paragraph.add_run()
        run.font.size = size
        run.font.color.rgb = color
        if part in ("begin", "separate", "end"):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), part)
        elif part == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = part
        run._r.append(el)


def style_doc(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in (("Heading 1", 18, BLUE), ("Heading 2", 13, INK), ("Title", 32, INK)):
        st = doc.styles[name]
        # Drop theme fonts so the explicit font wins.
        rfonts = st.element.get_or_add_rPr().find(qn("w:rFonts"))
        if rfonts is not None:
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rfonts.attrib.pop(qn(attr), None)
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(4)

    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    section.header_distance = section.footer_distance = Inches(0.45)


def header_footer(doc, logo, watermark):
    section = doc.sections[0]
    # The cover shows the large logo, so its header and footer carry no logo text there.
    section.different_first_page_header_footer = True

    # The watermark sits on every page, cover included.
    floating(section.first_page_header.paragraphs[0], watermark, Inches(5.5), None)
    floating(section.header.paragraphs[0], watermark, Inches(5.5), None)

    head = section.header.paragraphs[0]
    head.add_run().add_picture(io.BytesIO(logo), width=Inches(1.1))

    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = foot.add_run("{{ vendor_name }} · Scope of Work · {{ project_title }} · Confidential · Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    field(foot, "PAGE", Pt(8), MUTED)


def para(doc, text="", *, bold=False, size=None, color=None, align=None, after=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def h1(doc, number):
    return doc.add_heading(f"{number}. {SECTIONS[number - 1]}", level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def tag(doc, text):
    """A paragraph holding only a {%p %} tag; docxtpl removes it on render."""
    doc.add_paragraph(text)


def bullets(doc, var, text="{{ item }}"):
    """A {%p for item in var %} bullet list."""
    tag(doc, f"{{%p for item in {var} %}}")
    doc.add_paragraph(text, style="List Bullet")
    tag(doc, "{%p endfor %}")


def header_row(table, headers):
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, HEADER_FILL)
    # Repeat the header row when a table runs onto the next page.
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_widths(table, widths):
    """Fixed column widths; LibreOffice reads the grid, Word the cells, so set both."""
    table.autofit = False
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Inches(w)


def loop_table(doc, headers, loop, cells, total=None, widths=None):
    """Header row, then {%tr for %} / data / {%tr endfor %} rows, then an optional total row."""
    # docxtpl drops the two tag rows and repeats the data row.
    table = doc.add_table(rows=5 if total else 4, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_row(table, headers)
    table.rows[1].cells[0].text = f"{{%tr for {loop} %}}"
    for i, text in enumerate(cells):
        table.rows[2].cells[i].text = text
    table.rows[3].cells[0].text = "{%tr endfor %}"
    if total:
        for i, text in enumerate(total):
            table.rows[4].cells[i].paragraphs[0].add_run(text).bold = True
    set_widths(table, widths or [6.5 / len(headers)] * len(headers))
    doc.add_paragraph()
    return table


def optional_table(doc, var, empty_text, headers, loop, cells, widths=None):
    tag(doc, f"{{%p if {var} %}}")
    loop_table(doc, headers, loop, cells, widths=widths)
    tag(doc, "{%p else %}")
    doc.add_paragraph(empty_text)
    tag(doc, "{%p endif %}")


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def dashed_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "dashed")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "A6ACB0")
        borders.append(el)
    tc_pr.append(borders)


def logos(doc, logo):
    """Vendor logo and a slot for the client's, side by side."""
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    left, right = table.rows[0].cells
    left.paragraphs[0].add_run().add_picture(io.BytesIO(logo), width=Inches(2.2))
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    dashed_border(right)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(28)
    p.add_run(
        "{% if client_logo_path %}{{ client_logo }}{% else %}Customer logo{% endif %}"
    ).font.color.rgb = MUTED
    set_widths(table, [3.0, 3.0])
    return table


def cover(doc, logo):
    logos(doc, logo)
    for _ in range(4):
        para(doc)
    para(doc, "Scope of Work", bold=True, size=34, color=INK, after=4)
    para(doc, "{{ project_title }}", size=18, color=BLUE, after=36)

    rows = [
        ("Prepared for", "{{ client_name }}"),
        ("Prepared by", "{{ vendor_name }}"),
        ("Date", "{{ sow_date }}"),
        ("Version", "{{ version }}"),
        ("{%tr if author %}", None),
        ("Author", "{{ author }}"),
        ("{%tr endif %}", None),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    for r, (label, value) in enumerate(rows):
        cells = table.rows[r].cells
        if value is None:
            cells[0].text = label
            continue
        lab = cells[0].paragraphs[0].add_run(label)
        lab.font.color.rgb = MUTED
        lab.font.size = Pt(10)
        cells[1].paragraphs[0].add_run(value).bold = True
    set_widths(table, [1.5, 5.0])
    for _ in range(4):
        para(doc)
    para(
        doc,
        "This document is confidential and proprietary to the Vendor and the Client named "
        "above. It may not be disclosed to third parties without the prior written consent "
        "of both parties.",
        size=8,
        color=MUTED,
    )
    page_break(doc)


def preamble(doc):
    para(doc, "Scope of Work", bold=True, size=22, color=INK, after=0)
    para(doc, "{{ project_title }}", size=14, color=BLUE, after=12)
    doc.add_paragraph(
        "This Scope of Work (“SOW”) is agreed upon and entered into by and between "
        "{{ vendor_name }}, of {{ vendor_address }} (the “Vendor”), and {{ client_name }}"
        "{% if client_address %}, of {{ client_address }}{% endif %} (the “Client”). "
        "It is effective on the date of the last signature below. "
        "{% if msa_reference %}It is issued under, and governed by the terms of, the "
        "{{ msa_reference }} between the parties; where the two conflict, that agreement "
        "prevails.{% else %}It is governed by its own terms and supersedes any other agreement "
        "between the parties on its subject matter.{% endif %}"
        "{% if governing_law %} This SOW is governed by the laws of {{ governing_law }}.{% endif %}"
    )
    para(doc, "Contents", bold=True, size=13, color=INK, after=4).paragraph_format.space_before = Pt(18)
    for i, name in enumerate(SECTIONS, 1):
        para(doc, f"{i}.  {name}", after=2)
    page_break(doc)


def build(out: Path):
    doc = Document()
    style_doc(doc)
    logo = logo_png()
    header_footer(doc, logo, watermark_png())
    cover(doc, logo)
    preamble(doc)

    h1(doc, 1)
    doc.add_paragraph("{{ executive_summary }}")
    h2(doc, "Project overview")
    doc.add_paragraph(
        "This project consists of the following high-level phases, described in section 3:"
    )
    bullets(doc, "phases", text="{{ item.name }}")
    h2(doc, "Project objectives")
    doc.add_paragraph("The key objectives of this project are:")
    bullets(doc, "objectives")
    h2(doc, "1.1 Vendor overview")
    doc.add_paragraph("{{ vendor_overview }}")
    h2(doc, "1.2 Client overview")
    doc.add_paragraph("{{ client_overview }}")

    h1(doc, 2)
    h2(doc, "2.1 Functional requirements")
    doc.add_paragraph(
        "The project addresses the following functional requirements. They guide the "
        "current implementation."
    )
    loop_table(
        doc, ["#", "Functional requirement"], "r in functional_requirements",
        ["{{ loop.index }}", "{{ r }}"], widths=[0.5, 6.0],
    )
    h2(doc, "2.2 Non-functional requirements")
    tag(doc, "{%p if non_functional_requirements %}")
    doc.add_paragraph(
        "The project addresses the following non-functional requirements. They guide the "
        "current implementation."
    )
    loop_table(
        doc, ["#", "Non-functional requirement"], "r in non_functional_requirements",
        ["{{ loop.index }}", "{{ r }}"], widths=[0.5, 6.0],
    )
    tag(doc, "{%p else %}")
    doc.add_paragraph("No non-functional requirements specified.")
    tag(doc, "{%p endif %}")
    h2(doc, "2.3 Architecture overview")
    tag(doc, "{%p if architecture_overview %}")
    doc.add_paragraph("{{ architecture_overview }}")
    tag(doc, "{%p else %}")
    doc.add_paragraph(
        "The deployment architecture will be agreed between the Vendor and the Client and "
        "documented in a Technical Design Document."
    )
    tag(doc, "{%p endif %}")
    tag(doc, "{%p if architecture_components %}")
    doc.add_paragraph("The core components of the solution are:")
    loop_table(
        doc, ["Component", "Role in the solution"], "c in architecture_components",
        ["{{ c.name }}", "{{ c.description }}"], widths=[2.0, 4.5],
    )
    doc.add_paragraph("The list of components may change as the design is finalised.")
    tag(doc, "{%p endif %}")
    tag(doc, "{%p if integrations %}")
    doc.add_paragraph("The following integrations with the Client’s existing environment are in scope:")
    loop_table(
        doc, ["Integration", "Source, target and purpose"], "i in integrations",
        ["{{ i.name }}", "{{ i.description }}"], widths=[2.0, 4.5],
    )
    tag(doc, "{%p endif %}")

    h1(doc, 3)
    doc.add_paragraph("This project is organised into the phases described below.")
    tag(doc, "{%p for p in phases %}")
    h2(doc, "3.{{ loop.index }} {{ p.name }}")
    doc.add_paragraph("{{ p.focus }}")
    doc.add_paragraph("Key tasks in this phase:")
    tag(doc, "{%p for t in p.tasks %}")
    doc.add_paragraph("{{ t }}", style="List Bullet")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p endfor %}")

    h1(doc, 4)
    doc.add_paragraph("Deliverables for this engagement are:")
    loop_table(
        doc,
        ["Phase", "Deliverable", "Description", "Format"],
        "d in deliverables",
        ["{{ d.phase }}", "{{ d.name }}", "{{ d.description }}", "{{ d.format }}"],
        widths=[1.4, 1.6, 2.4, 1.1],
    )
    doc.add_paragraph("{{ acceptance_criteria }}")

    h1(doc, 5)
    h2(doc, "5.1 Out of scope")
    tag(doc, "{%p if out_of_scope %}")
    doc.add_paragraph("The following are not in scope for this SOW:")
    bullets(doc, "out_of_scope")
    tag(doc, "{%p else %}")
    doc.add_paragraph("No exclusions specified.")
    tag(doc, "{%p endif %}")
    h2(doc, "5.2 Assumptions")
    tag(doc, "{%p if assumptions %}")
    doc.add_paragraph("This SOW, its timeline and its fees rely on the following assumptions:")
    bullets(doc, "assumptions")
    tag(doc, "{%p else %}")
    doc.add_paragraph("No assumptions specified.")
    tag(doc, "{%p endif %}")
    h2(doc, "5.3 Risks")
    tag(doc, "{%p if risks %}")
    doc.add_paragraph(
        "The following shared risks are managed jointly by the Vendor and the Client "
        "throughout the project."
    )
    loop_table(
        doc, ["Risk", "Mitigation"], "k in risks",
        ["{{ k.risk }}", "{{ k.mitigation }}"], widths=[2.5, 4.0],
    )
    tag(doc, "{%p else %}")
    doc.add_paragraph("No specific risks identified at the time of writing.")
    tag(doc, "{%p endif %}")
    h2(doc, "5.4 Change control")
    doc.add_paragraph("{{ change_control }}")

    h1(doc, 6)
    doc.add_paragraph(
        "The success criteria below summarise the overall outcomes and set clear expectations "
        "for the completion of this SOW’s activities and deliverables."
    )
    doc.add_paragraph("Completion of all deliverables documented in this SOW", style="List Bullet")
    bullets(doc, "success_criteria")

    h1(doc, 7)
    doc.add_paragraph(
        "The following is an estimated timeline. Actual timelines may vary and will be "
        "reviewed regularly with the Client."
    )
    p = doc.add_paragraph()
    p.add_run("Start date (estimated): ").bold = True
    p.add_run("{{ start_date }}")
    p = doc.add_paragraph()
    p.add_run("End date (estimated): ").bold = True
    p.add_run("{{ end_date }}, pending acceptance of all deliverables")
    loop_table(
        doc,
        ["Phase", "Period", "Activities and outcomes"],
        "t in timeline",
        ["{{ t.phase }}", "{{ t.period }}", "{{ t.activities }}"],
        widths=[1.7, 1.2, 3.6],
    )

    h1(doc, 8)
    doc.add_paragraph("The Vendor will provide the following roles for this engagement.")
    loop_table(
        doc, ["Vendor role", "Responsibilities"], "r in vendor_roles",
        ["{{ r.role }}", "{{ r.responsibilities }}"], widths=[2.2, 4.3],
    )
    tag(doc, "{%p if client_roles %}")
    doc.add_paragraph(
        "The Client will provide the following roles. They need not be dedicated full time, "
        "but must take part in the relevant reviews, workshops and tasks."
    )
    loop_table(
        doc, ["Client role", "Responsibilities"], "r in client_roles",
        ["{{ r.role }}", "{{ r.responsibilities }}"], widths=[2.2, 4.3],
    )
    tag(doc, "{%p else %}")
    doc.add_paragraph("The Client’s project roles will be confirmed at kick-off.")
    tag(doc, "{%p endif %}")

    h1(doc, 9)
    h2(doc, "9.1 Pricing")
    doc.add_paragraph("Pricing model: {{ pricing_model }}. All amounts are in {{ currency }}.")
    subtotal_label = "{% if fees %}Subtotal{% else %}Total cost{% endif %}"
    tag(doc, "{%p if project_type == 'package' %}")
    doc.add_paragraph("The package comprises:")
    loop_table(
        doc,
        ["Item", "Description", "Qty", "Unit price", "Amount"],
        "i in package_items",
        ["{{ i.item }}", "{{ i.description }}", "{{ i.quantity }}", "{{ i.unit_price }}", "{{ i.amount }}"],
        total=[subtotal_label, "", "", "", "{{ subtotal }}"],
        widths=[1.6, 2.0, 0.6, 1.15, 1.15],
    )
    tag(doc, "{%p else %}")
    doc.add_paragraph("Professional services are estimated by role in mandays:")
    loop_table(
        doc,
        ["Role", "Mandays", "Day rate", "Amount"],
        "e in effort",
        ["{{ e.role }}", "{{ e.mandays }}", "{{ e.day_rate }}", "{{ e.amount }}"],
        total=[subtotal_label, "{{ total_mandays }}", "", "{{ subtotal }}"],
        widths=[2.6, 1.0, 1.4, 1.5],
    )
    tag(doc, "{%p endif %}")
    tag(doc, "{%p if fees %}")
    doc.add_paragraph("Other fees and adjustments:")
    loop_table(
        doc, ["Item", "Amount"], "f in fees",
        ["{{ f.item }}", "{{ f.amount }}"],
        total=["Total cost", "{{ total_fees }}"], widths=[5.0, 1.5],
    )
    tag(doc, "{%p endif %}")
    doc.add_paragraph(
        "Payments are invoiced on the following milestones. Each milestone’s deliverables "
        "must be completed and accepted by the Client before it is invoiced."
    )
    loop_table(
        doc,
        ["Milestone", "Deliverable(s)", "Estimated completion", "Payment"],
        "p in payment_schedule",
        ["{{ p.milestone }}", "{{ p.deliverables }}", "{{ p.due }}", "{{ p.amount }}"],
        widths=[1.5, 2.4, 1.3, 1.3],
    )
    doc.add_paragraph("Payment terms: {{ payment_terms }}.")
    h2(doc, "9.2 Expenses")
    doc.add_paragraph("{{ expenses }}")
    h2(doc, "9.3 Taxes")
    doc.add_paragraph(
        "{% if tax_treatment == 'inclusive' %}Pricing above is inclusive of any applicable "
        "sales tax or VAT. The Client shall provide a valid tax exemption certificate if "
        "required.{% else %}Pricing above is exclusive of any applicable sales tax or VAT. "
        "The Client is liable for and shall pay all sales tax or VAT properly payable in "
        "connection with the services under this SOW, and shall provide a valid tax "
        "exemption certificate if required.{% endif %}"
    )

    h1(doc, 10)
    doc.add_paragraph("Signed by the authorised representatives of the parties.")
    sig = doc.add_table(rows=5, cols=2)
    sig.style = "Table Grid"
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(sig, [3.25, 3.25])
    rows = [
        ("For {{ client_name }}", "For {{ vendor_name }}"),
        ("Name: {{ client_signatory_name }}", "Name: {{ vendor_signatory_name }}"),
        ("Title: {{ client_signatory_title }}", "Title: {{ vendor_signatory_title }}"),
        ("Signature:\n\n", "Signature:\n\n"),
        ("Date:", "Date:"),
    ]
    for r, (left, right) in enumerate(rows):
        sig.rows[r].cells[0].text = left
        sig.rows[r].cells[1].text = right
    for cell in sig.rows[0].cells:
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, HEADER_FILL)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    print(build(parser.parse_args().out))
